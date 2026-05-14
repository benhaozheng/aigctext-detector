# -*- coding: utf-8 -*-
"""
文件读取模块
功能：支持txt/pdf/docx文件读取
"""

import os
from typing import Optional
import warnings
warnings.filterwarnings('ignore')


class FileReader:
    """文件读取器"""

    @staticmethod
    def read_file(file_path: str, encoding: str = 'utf-8') -> Optional[str]:
        """
        读取文件内容（自动检测文件类型）

        Args:
            file_path: 文件路径
            encoding: 文本文件编码

        Returns:
            文件内容字符串，读取失败返回None
        """
        if not os.path.exists(file_path):
            print(f"文件不存在: {file_path}")
            return None

        # 获取文件扩展名
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()

        try:
            # 根据文件类型选择读取方法
            if ext == '.txt':
                return FileReader._read_txt(file_path, encoding)
            elif ext == '.pdf':
                return FileReader._read_pdf(file_path)
            elif ext == '.docx':
                return FileReader._read_docx(file_path)
            else:
                # 尝试作为文本文件读取
                print(f"未知文件类型: {ext}，尝试作为文本文件读取")
                return FileReader._read_txt(file_path, encoding)
        except Exception as e:
            print(f"读取文件失败: {e}")
            return None

    @staticmethod
    def _read_txt(file_path: str, encoding: str = 'utf-8') -> str:
        """读取txt文件"""
        # 尝试多种编码
        encodings = [encoding, 'utf-8', 'gbk', 'gb2312', 'utf-8-sig']

        for enc in encodings:
            try:
                with open(file_path, 'r', encoding=enc) as f:
                    content = f.read()
                print(f"✅ 成功读取txt文件（编码: {enc}）")
                return content
            except UnicodeDecodeError:
                continue

        raise Exception("无法解码文件，请指定正确的编码")

    @staticmethod
    def _read_pdf(file_path: str) -> str:
        """读取pdf文件"""
        try:
            import PyPDF2
        except ImportError:
            print("⚠️ 需要安装 PyPDF2: pip install PyPDF2")
            raise

        text = []
        with open(file_path, 'rb') as f:
            pdf_reader = PyPDF2.PdfReader(f)

            # 获取页数
            num_pages = len(pdf_reader.pages)
            print(f"PDF文件共 {num_pages} 页")

            # 逐页提取文本
            for page_num in range(num_pages):
                page = pdf_reader.pages[page_num]
                page_text = page.extract_text()
                text.append(page_text)

        content = '\n\n'.join(text)
        print(f"✅ 成功读取pdf文件")
        return content

    @staticmethod
    def _read_docx(file_path: str) -> str:
        """读取docx文件"""
        try:
            from docx import Document
        except ImportError:
            print("⚠️ 需要安装 python-docx: pip install python-docx")
            raise

        doc = Document(file_path)

        # 提取段落文本
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())

        # 提取表格文本
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(' | '.join(row_text))

        content = '\n\n'.join(paragraphs)
        print(f"✅ 成功读取docx文件（共{len(paragraphs)}个段落）")
        return content

    @staticmethod
    def read_multiple_files(file_paths: list) -> dict:
        """
        批量读取多个文件

        Args:
            file_paths: 文件路径列表

        Returns:
            {文件路径: 文件内容} 字典
        """
        results = {}

        for file_path in file_paths:
            content = FileReader.read_file(file_path)
            if content:
                results[file_path] = content

        return results

    @staticmethod
    def read_directory(directory: str, extensions: list = None) -> dict:
        """
        读取目录下所有文件

        Args:
            directory: 目录路径
            extensions: 文件扩展名列表，如['.txt', '.pdf']

        Returns:
            {文件路径: 文件内容} 字典
        """
        if extensions is None:
            extensions = ['.txt', '.pdf', '.docx']

        results = {}

        for filename in os.listdir(directory):
            file_path = os.path.join(directory, filename)

            if os.path.isfile(file_path):
                _, ext = os.path.splitext(filename)
                if ext.lower() in extensions:
                    content = FileReader.read_file(file_path)
                    if content:
                        results[file_path] = content

        return results


if __name__ == '__main__':
    # 测试代码
    reader = FileReader()

    print("=" * 50)
    print("文件读取器测试")
    print("=" * 50)

    # 测试txt文件
    test_txt = "./data/test.txt"
    if os.path.exists(test_txt):
        print(f"\n测试TXT文件: {test_txt}")
        content = reader.read_file(test_txt)
        print(f"内容长度: {len(content)} 字符")
        print(f"内容预览: {content[:100]}...")

    # 测试pdf文件（如果有）
    test_pdf = "./data/test.pdf"
    if os.path.exists(test_pdf):
        print(f"\n测试PDF文件: {test_pdf}")
        content = reader.read_file(test_pdf)
        print(f"内容长度: {len(content)} 字符")
        print(f"内容预览: {content[:100]}...")

    # 测试docx文件（如果有）
    test_docx = "./data/test.docx"
    if os.path.exists(test_docx):
        print(f"\n测试DOCX文件: {test_docx}")
        content = reader.read_file(test_docx)
        print(f"内容长度: {len(content)} 字符")
        print(f"内容预览: {content[:100]}...")
